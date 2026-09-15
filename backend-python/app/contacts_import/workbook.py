"""Reads AI TEAM - CURRENT CONTACTS.xlsx and produces per-row outcomes.

The only module here that touches openpyxl. Forward-fill for JOB AGENCIES UK
lives here (not in rows.py) because it's a cross-row concern: which company
a contact/job row belongs to depends on rows above it, which is exactly the
kind of state rows.py's pure, single-row functions are deliberately kept
free of.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from app.contacts_import.rows import (
    AgencyRowOutcome,
    CompanyRowOutcome,
    RawCell,
    extract_agency_row,
    extract_company_name,
    extract_company_row,
)

JOB_AGENCIES_SHEET = "JOB AGENCIES UK"
UK_COMPANIES_SHEET = "UK COMPANIES"


def _cell(ws: Worksheet, row: int, col: int) -> RawCell:
    cell = ws.cell(row=row, column=col)
    hyperlink = cell.hyperlink.target if cell.hyperlink else None
    return RawCell(value=cell.value, hyperlink=hyperlink)


def load_workbook(path: Path) -> openpyxl.Workbook:
    # data_only=True: read cached formula results, not formula text — none of
    # the columns here are formulas in the source file, but this is the safe
    # default for a workbook that might be re-exported from Google Sheets.
    return openpyxl.load_workbook(path, data_only=True)


def read_job_agencies_uk(wb: openpyxl.Workbook) -> list[AgencyRowOutcome]:
    """Company, Link Job/Company, LinkedIn, Phone, Email — with column A
    (Company) forward-filled: a blank cell means "same company as the row
    above", not "no company"."""
    ws = wb[JOB_AGENCIES_SHEET]

    outcomes: list[AgencyRowOutcome] = []
    current_company = None
    have_seen_company = False

    for row in range(2, ws.max_row + 1):
        company_cell = _cell(ws, row, 1)
        link_cell = _cell(ws, row, 2)
        linkedin_cell = _cell(ws, row, 3)
        phone_cell = _cell(ws, row, 4)
        email_cell = _cell(ws, row, 5)

        if all(
            c.value is None
            for c in (company_cell, link_cell, linkedin_cell, phone_cell, email_cell)
        ):
            # A fully blank row (spacer between agency groups) carries
            # nothing to forward-fill or extract.
            continue

        if company_cell.value is not None:
            current_company = extract_company_name(
                company_cell.value, company_cell.hyperlink, allow_domain_fallback=True
            )
            have_seen_company = True
        elif not have_seen_company:
            current_company = extract_company_name(
                None, None, allow_domain_fallback=True
            )

        outcomes.append(
            extract_agency_row(
                row,
                current_company,
                link_cell,
                linkedin_cell,
                phone_cell.value,
                email_cell.value,
            )
        )

    return outcomes


def read_uk_companies(wb: openpyxl.Workbook) -> list[CompanyRowOutcome]:
    """COMPANY, JOB TITLE/LINK — no forward-fill: every row with data names
    its own company (see the importer's mapping-decisions report for the
    audit confirming this)."""
    ws = wb[UK_COMPANIES_SHEET]

    outcomes: list[CompanyRowOutcome] = []
    for row in range(2, ws.max_row + 1):
        company_cell = _cell(ws, row, 1)
        job_cell = _cell(ws, row, 2)

        if company_cell.value is None and job_cell.value is None:
            continue

        company = extract_company_name(
            company_cell.value, company_cell.hyperlink, allow_domain_fallback=False
        )
        outcomes.append(extract_company_row(row, company, job_cell))

    return outcomes


def read_all(path: Path) -> tuple[list[AgencyRowOutcome], list[CompanyRowOutcome]]:
    wb = load_workbook(path)
    missing = {JOB_AGENCIES_SHEET, UK_COMPANIES_SHEET} - set(wb.sheetnames)
    if missing:
        raise ValueError(
            f"Workbook is missing expected sheet(s): {sorted(missing)}. "
            f"Found: {wb.sheetnames}"
        )
    return read_job_agencies_uk(wb), read_uk_companies(wb)
